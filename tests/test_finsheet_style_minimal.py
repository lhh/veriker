"""tests/test_finsheet_style_minimal.py — the FinSheet-style pilot and its two
promoted primitives.

Three things, each without a tautology:

  FAITHFULNESS. The generator's gold (openpyxl records, its own Decimal
  arithmetic, no audit_bundle import) must agree with the verifier primitive's
  recompute (stdlib zip+xml parse of the SAME workbook bytes) on every question
  of a layout sample. The two implementations share no code; agreement is a
  finding, and the first run of this check found a real one-cent disagreement
  (openpyxl's 16-significant-digit float serialisation, fixed producer-side).

  TEETH PER ARM. The five demo profiles go through the pilot's anchored
  verifier — the same `auditor_entry.build_verifier` the harness uses — and
  each lands on the arm the pre-registration says it should, INCLUDING the
  wrong_selection profile that Arm B is documented NOT to catch. A test that
  only proved catches would let that scope limit drift into a claim.

  PIN. The committed demo query bytes equal what the generator produces for
  the demo question, so a generator change that moves the question fails here
  rather than silently making the committed spec pin a question nobody asks.

Run in one process; the demo builds are ~1s each.
"""

from __future__ import annotations

import importlib.util
import json
import os
import sys
from pathlib import Path

import pytest

_PKG_ROOT = Path(__file__).resolve().parents[1]
if str(_PKG_ROOT) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT))
_PILOT = _PKG_ROOT / "examples" / "finsheet_style_minimal"
if str(_PILOT) not in sys.path:
    sys.path.insert(0, str(_PILOT))

from audit_bundle.rederivation.primitives.sheet_query import (  # noqa: E402
    QueryRefused,
    evaluate_query,
    read_workbook,
    replay_derivation,
)


def _load(name: str):
    spec = importlib.util.spec_from_file_location(
        f"finsheet_style_minimal__{name}", _PILOT / f"{name}.py"
    )
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = mod  # @dataclass needs the module registered
    spec.loader.exec_module(mod)
    return mod


# MUST precede the _load calls below: _generate_set.py imports openpyxl at its
# own module level, so a guard placed after it never runs — the module raises
# ModuleNotFoundError during collection and takes the whole file down with it.
# (It sat one line lower until 2026-09-03; a lean install turned this file into a
# collection ERROR instead of the clean SKIP the guard was written to produce.)
# openpyxl is a GENERATOR dependency only — the verifier never needs it.
pytest.importorskip("openpyxl")

# Loaded by PATH under pilot-unique names: a bare `import auditor_entry` in a
# shared pytest process resolves to whichever pilot's module was imported
# first (corner_load_equilibrium_minimal ships one too) — measured 2026-09-02,
# ten of these tests failed in the full suite and passed solo.
ae = _load("auditor_entry")
_gen = _load("_generate_set")
_build = _load("_build_bundle")


# ---------------------------------------------------------------------------
# Faithfulness: generator gold == primitive recompute, one file per layout
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "file_id", ["L1_S1", "L2_S2", "L3_S2", "L4_S1", "L5_S2", "L6_S2"]
)
def test_generator_gold_matches_primitive_recompute(tmp_path, file_id):
    wb_path, _recs, qs = _gen.build_one(file_id, tmp_path)
    book = read_workbook(wb_path)
    assert len(qs) == 16
    for q in qs:
        got = evaluate_query(book, q["query"])
        gold = q["gold"]
        if q["answer_kind"] in ("number",):
            assert abs(got - gold) <= 1e-9, (file_id, q["qid"], got, gold)
        else:
            assert got == gold, (file_id, q["qid"], got, gold)


def test_mutant_gold_is_a_disagreement(tmp_path):
    """Control for the test above: a one-cent edit to the generator's gold
    must NOT agree. Without this, 'agreement' could be vacuous."""
    wb_path, _recs, qs = _gen.build_one("L2_S1", tmp_path)
    book = read_workbook(wb_path)
    q = next(x for x in qs if x["qid"] == "Q10")
    got = evaluate_query(book, q["query"])
    assert abs(got - (q["gold"] + 0.01)) > 1e-9


def test_workbook_numbers_are_canonical(tmp_path):
    """The producer-side canonicalisation holds: no cell carries more than two
    decimals (the openpyxl 16-significant-digit artefact is gone)."""
    import re
    import zipfile

    wb_path, _recs, _qs = _gen.build_one("L5_S2", tmp_path)
    z = zipfile.ZipFile(wb_path)
    long_cells = 0
    for n in z.namelist():
        if n.startswith("xl/worksheets/"):
            long_cells += len(re.findall(rb"<v>-?\d+\.\d{3,}</v>", z.read(n)))
    assert long_cells == 0


# ---------------------------------------------------------------------------
# The committed demo question is pinned to what the generator produces
# ---------------------------------------------------------------------------


def test_committed_demo_query_matches_generator(tmp_path):
    _wb, _recs, qs = _gen.build_one(_build.DEMO_FILE_ID, tmp_path)
    q = next(x for x in qs if x["qid"] == _build.DEMO_QID)
    expected = json.dumps(q["query"], indent=2, sort_keys=True).encode("utf-8")
    assert ae.DEMO_QUERY.read_bytes() == expected, (
        "regenerate spec_pinned/demo_query.json + the demo spec"
    )
    spec = json.loads(ae.DEMO_SPEC.read_bytes())
    assert spec["types"][ae.OUTPUT_SEMANTIC]["pinned_inputs"][
        ae.QUERY_REL
    ] == ae.sha256_bytes(expected)


# ---------------------------------------------------------------------------
# Teeth per arm, through the pilot's own anchored verifier
# ---------------------------------------------------------------------------


def _run_profile(tmp_path, profile: str) -> dict:
    bundle = _build.build_demo(tmp_path / f"bundle_{profile}", profile)
    verifier, _anchor = ae.build_verifier(bundle, ae.DEMO_SPEC)
    return ae.read_outcomes(verifier.verify(bundle))


def test_clean_passes_both_arms(tmp_path):
    out = _run_profile(tmp_path, "clean")
    assert out["exit_code"] == 0 and out["ok"], out
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["state"] == "RE_DERIVED"
    assert out["outputs"][ae.OUTPUT_CONSISTENCY]["state"] == "RE_DERIVED"


def test_wrong_arithmetic_fails_both_arms(tmp_path):
    out = _run_profile(tmp_path, "wrong_arithmetic")
    assert out["exit_code"] == 1
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["reason_code"] == "RE_DERIVATION_MISMATCH"
    assert (
        out["outputs"][ae.OUTPUT_CONSISTENCY]["reason_code"] == "RE_DERIVATION_MISMATCH"
    )


def test_wrong_selection_is_caught_by_A_and_NOT_by_B(tmp_path):
    """The documented scope limit of sheet_derivation_replay, demonstrated:
    the wrong column summed correctly replays to its own number."""
    out = _run_profile(tmp_path, "wrong_selection")
    assert out["exit_code"] == 1
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["reason_code"] == "RE_DERIVATION_MISMATCH"
    assert out["outputs"][ae.OUTPUT_CONSISTENCY]["state"] == "RE_DERIVED"


def test_hallucinated_operand_is_named_by_B(tmp_path):
    out = _run_profile(tmp_path, "hallucinated_operand")
    assert out["exit_code"] == 1
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["reason_code"] == "RE_DERIVATION_MISMATCH"
    row = out["outputs"][ae.OUTPUT_CONSISTENCY]
    assert row["reason_code"] == "RE_DERIVATION_MISMATCH"
    assert "operand_not_in_sheet" in row["detail"]


def test_wrong_question_trips_the_pin(tmp_path):
    """Answering an easier question correctly and shipping it as the auditor's
    is refused on the pin, before any recompute."""
    out = _run_profile(tmp_path, "wrong_question")
    assert out["exit_code"] == 1
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["reason_code"] == "PINNED_INPUT_MISMATCH"


def test_anchor_from_inside_bundle_is_refused(tmp_path):
    from audit_bundle.rederivation.spec_binding import AnchorConstructionError

    bundle = _build.build_demo(tmp_path / "bundle_clean", "clean")
    with pytest.raises(AnchorConstructionError):
        ae.build_verifier(bundle, bundle / "spec" / ae.DEMO_SPEC.name)


def test_deleting_manifest_outputs_is_not_a_pass(tmp_path):
    """require_rederivation + work-set: dropping the declaration is a finding."""
    bundle = _build.build_demo(tmp_path / "bundle_clean", "clean")
    mp = bundle / "manifest.json"
    m = json.loads(mp.read_text())
    m.pop("outputs")
    mp.write_text(json.dumps(m, indent=2, sort_keys=True))
    verifier, _ = ae.build_verifier(bundle, ae.DEMO_SPEC)
    out = ae.read_outcomes(verifier.verify(bundle))
    assert out["exit_code"] != 0


# ---------------------------------------------------------------------------
# Primitive-level edge cases (no bundle)
# ---------------------------------------------------------------------------


def test_lookup_ambiguity_is_a_refusal(tmp_path):
    wb_path, _recs, qs = _gen.build_one("L1_S1", tmp_path)
    book = read_workbook(wb_path)
    q = dict(next(x for x in qs if x["qid"] == "Q01")["query"])
    q = {**q, "filters": []}  # every row matches
    with pytest.raises(QueryRefused):
        evaluate_query(book, q)


def test_replay_refuses_malformed_and_names_missing_operands(tmp_path):
    wb_path, recs, _qs = _gen.build_one("L1_S1", tmp_path)
    book = read_workbook(wb_path)
    assert replay_derivation(book, {"schema": "nope"})["kind"] == "refused"
    nav = float(recs[0]["nav"])
    ok = replay_derivation(
        book, {"schema": "sheet-derivation-v1", "op": "sum", "operands": [nav]}
    )
    assert ok == round(nav, 2)
    bad = replay_derivation(
        book, {"schema": "sheet-derivation-v1", "op": "sum", "operands": [nav, 0.123]}
    )
    assert bad["kind"] == "operand_not_in_sheet" and bad["missing"] == [0.123]


def test_reader_refuses_dtd(tmp_path):
    import zipfile

    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _recs, _qs = _gen.build_one("L1_S1", tmp_path)
    evil = tmp_path / "evil.xlsx"
    src = zipfile.ZipFile(wb_path)
    with zipfile.ZipFile(evil, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            # openpyxl writes inline strings, so there is no sharedStrings.xml
            # to poison; workbook.xml is always present and always parsed.
            if info.filename == "xl/workbook.xml":
                data = b'<!DOCTYPE x [<!ENTITY a "b">]>' + data
            dst.writestr(info.filename, data)
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


# ---------------------------------------------------------------------------
# Adversarial pass 2026-09-02: one test per finding CLASS, not per instance
# ---------------------------------------------------------------------------


def _clean_bundle(tmp_path):
    return _build.build_demo(tmp_path / "bundle_clean", "clean")


def _verify(bundle, spec=None):
    verifier, _ = ae.build_verifier(bundle, spec or ae.DEMO_SPEC)
    return ae.read_outcomes(verifier.verify(bundle))


def test_workbook_substitution_trips_the_pin(tmp_path):
    """CRITICAL finding: a producer that ships the workbook under which its
    wrong answer is right, re-cohering the manifest, passed both arms. The
    workbook is now pinned on both types."""
    import openpyxl

    bundle = _clean_bundle(tmp_path)
    wb_path = bundle / ae.WORKBOOK_REL
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Holdings"]
    # bump the first numeric NAV cell under the header by 1.00
    hdr = next(r for r in range(1, 12) if ws.cell(r, 1).value == "Company")
    for r in range(hdr + 1, ws.max_row + 1):
        v = ws.cell(r, 7).value
        if isinstance(v, (int, float)) and ws.cell(r, 1).value not in (
            "Subtotal",
            "Total",
        ):
            ws.cell(r, 7).value = float(v) + 1.0
            break
    wb.save(wb_path)
    _build.recohere_file(bundle, ae.WORKBOOK_REL)
    claim = json.dumps({"value": 589.72}).encode()
    for oid in (ae.OUTPUT_SEMANTIC, ae.OUTPUT_CONSISTENCY):
        (bundle / f"outputs/{oid}.json").write_bytes(claim)
        _build.recohere_file(bundle, f"outputs/{oid}.json")
    out = _verify(bundle)
    assert out["exit_code"] == 1
    assert out["outputs"][ae.OUTPUT_SEMANTIC]["reason_code"] == "PINNED_INPUT_MISMATCH"
    assert (
        out["outputs"][ae.OUTPUT_CONSISTENCY]["reason_code"] == "PINNED_INPUT_MISMATCH"
    )


def test_refusal_value_cannot_be_claimed(tmp_path):
    """SEVERE finding: a refusal was a plain dict, so a producer who claimed
    that same dict equalled it under `exact` (and its keys under `set`) with an
    empty derivation. Refusals now carry a per-call nonce. Tested on the count
    and list kinds, the two comparators that were open."""
    wb_path, recs, qs = _gen.build_one("L2_S1", tmp_path / "set")
    for qid, kind in (("Q07", "count"), ("Q05", "list")):
        q = next(x for x in qs if x["qid"] == qid)
        assert q["answer_kind"] == kind
        qb = json.dumps(q["query"], indent=2, sort_keys=True).encode()
        spec_dir = tmp_path / f"anchor_{qid}"
        spec_dir.mkdir()
        spec_path = spec_dir / "q.spec.json"
        spec_path.write_bytes(ae.make_spec(f"t.{qid}", kind, qb, wb_path.read_bytes()))
        # honest semantic claim; the CONSISTENCY claim is the refusal shape
        # the empty derivation will produce
        bundle = tmp_path / f"bundle_{qid}"
        _build.build_bundle(
            bundle,
            workbook_path=wb_path,
            query_bytes=qb,
            spec_bytes=spec_path.read_bytes(),
            spec_basename=spec_path.name,
            claimed_value=q["gold"],
            derivation={
                "schema": "sheet-derivation-v1",
                "op": "lookup",
                "operands": [],
            },
            bundle_id=f"t-{qid}",
        )
        forged = {"kind": "refused", "reason": "lookup needs exactly one operand"}
        if kind == "list":
            forged_v = ["kind", "reason", "nonce"]
        else:
            forged_v = forged
        (bundle / f"outputs/{ae.OUTPUT_CONSISTENCY}.json").write_bytes(
            json.dumps({"value": forged_v}).encode()
        )
        _build.recohere_file(bundle, f"outputs/{ae.OUTPUT_CONSISTENCY}.json")
        out = _verify(bundle, spec_path)
        assert out["outputs"][ae.OUTPUT_SEMANTIC]["state"] == "RE_DERIVED", (qid, out)
        assert out["outputs"][ae.OUTPUT_CONSISTENCY]["state"] == "NOT_RE_DERIVED", (
            qid,
            out,
        )
        assert out["exit_code"] == 1


def _rezip(src: Path, dst: Path, edit):
    import zipfile

    z = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w") as out:
        for info in z.infolist():
            out.writestr(info.filename, edit(info.filename, z.read(info.filename)))


def _inject_cell(xml: bytes, ref: bytes, v: bytes) -> bytes:
    return xml.replace(
        b"</sheetData>",
        b'<row><c r="' + ref + b'"><v>' + v + b"</v></c></row></sheetData>",
        1,
    )


@pytest.mark.parametrize("ref", [b"A5000000", b"ZZZ5", b"A0", b"XFE1"])
def test_reader_refuses_out_of_grid_references(tmp_path, ref):
    """SEVERE finding: one cell at row 20,000,000 materialised a 4 GB grid.
    Rows past 1,048,576, columns past 16,384, and row 0 are refused."""
    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    evil = tmp_path / "evil.xlsx"
    _rezip(
        wb_path,
        evil,
        lambda n, d: (
            _inject_cell(d, ref, b"1") if n == "xl/worksheets/sheet1.xml" else d
        ),
    )
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


def test_reader_refuses_grid_product(tmp_path):
    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    evil = tmp_path / "evil.xlsx"
    # 1,000,000 rows x 12 columns = 12M cells > cap, every reference in range
    _rezip(
        wb_path,
        evil,
        lambda n, d: (
            _inject_cell(d, b"A1000000", b"1") if n == "xl/worksheets/sheet1.xml" else d
        ),
    )
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


@pytest.mark.parametrize("pad", [b" " * 5000, b"<!--" + b"x" * 4200 + b"-->"])
def test_reader_refuses_doctype_anywhere(tmp_path, pad):
    """SEVERE finding: the guard scanned 4096 bytes; whitespace or a comment
    before <!DOCTYPE> got an entity expanded into a cell at exit 0."""
    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    evil = tmp_path / "evil.xlsx"

    def edit(n, d):
        if n == "xl/worksheets/sheet1.xml":
            body = d.split(b"?>", 1)[1] if d.startswith(b"<?xml") else d
            return pad + b'<!DOCTYPE worksheet [<!ENTITY inj "INJECTED">]>' + body
        return d

    _rezip(wb_path, evil, edit)
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


@pytest.mark.parametrize("v", [b"Infinity", b"NaN", b"sNaN", b"1e999999"])
def test_reader_refuses_non_finite_cells(tmp_path, v):
    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    evil = tmp_path / "evil.xlsx"
    _rezip(
        wb_path,
        evil,
        lambda n, d: (
            _inject_cell(d, b"Z2", v) if n == "xl/worksheets/sheet1.xml" else d
        ),
    )
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


@pytest.mark.parametrize("v", ["sNaN", "NaN", "Infinity", float("inf"), float("nan")])
def test_replay_non_finite_operand_is_absent_not_a_crash(tmp_path, v):
    """SEVERE finding: a 'sNaN' operand raised TypeError out of the hash set."""
    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    book = read_workbook(wb_path)
    out = replay_derivation(
        book, {"schema": "sheet-derivation-v1", "op": "sum", "operands": [v]}
    )
    assert isinstance(out, dict) and out["kind"] in (
        "operand_not_in_sheet",
        "refused",
    ), out
    assert "nonce" in out


def test_reader_refuses_corrupt_entry_bytes(tmp_path):
    """SEVERE finding: an entry whose stored bytes fail their CRC passed the
    declared-size caps and raised BadZipFile uncaught out of zf.read. The
    corruption must be on an entry the reader OPENS (xl/workbook.xml);
    [Content_Types].xml is never read, so corrupting it proves nothing."""
    import struct
    import zipfile

    from audit_bundle.rederivation.primitives.sheet_query import WorkbookUnreadable

    wb_path, _r, _q = _gen.build_one("L1_S1", tmp_path)
    good = tmp_path / "good.xlsx"
    _rezip(wb_path, good, lambda n, d: d)
    raw = bytearray(good.read_bytes())
    target = b"xl/workbook.xml"
    # local file header for the target
    pos = 0
    local = None
    while True:
        pos = raw.find(b"PK\x03\x04", pos)
        if pos < 0:
            break
        fnlen = struct.unpack("<H", raw[pos + 26 : pos + 28])[0]
        if raw[pos + 30 : pos + 30 + fnlen] == target:
            local = pos
            break
        pos += 4
    assert local is not None
    bad_crc = struct.pack(
        "<I", struct.unpack("<I", raw[local + 14 : local + 18])[0] ^ 0xFFFFFFFF
    )
    raw[local + 14 : local + 18] = bad_crc
    # central directory entry for the target
    pos = 0
    while True:
        pos = raw.find(b"PK\x01\x02", pos)
        if pos < 0:
            break
        fnlen = struct.unpack("<H", raw[pos + 28 : pos + 30])[0]
        if raw[pos + 46 : pos + 46 + fnlen] == target:
            raw[pos + 16 : pos + 20] = bad_crc
            break
        pos += 4
    evil = tmp_path / "evil.xlsx"
    evil.write_bytes(bytes(raw))
    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(evil).read(
            "xl/workbook.xml"
        )  # the raw failure the reader must map
    with pytest.raises(WorkbookUnreadable):
        read_workbook(evil)


@pytest.mark.parametrize("label", ["Subtotal ", "SUBTOTAL", " subtotal"])
def test_exclusion_and_divider_regexes_see_normalised_text(tmp_path, label):
    """MODERATE finding: 'Subtotal ' (trailing space) and 'SUBTOTAL' were summed
    as companies because the anchored regex saw the raw cell."""
    import openpyxl

    from audit_bundle.rederivation.primitives.sheet_query import extract_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(["Company", "MOIC (x)", "Unrealized Value ($m)"])
    ws.append(["fund i"])
    ws.append(["Acme", 1.5, 3.0])
    ws.append([label, None, 3.0])
    ws.append(["Total", None, 3.0])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    layout = {
        "sheet": "Holdings",
        "header_match": ["Company", "MOIC (x)"],
        "columns": {
            "company": "Company",
            "moic": "MOIC (x)",
            "nav": "Unrealized Value ($m)",
        },
        "group_from": r"divider:^Fund\s+[IVX]+$",
        "exclude_first_cell": "^(Total|Subtotal|Units)$",
        "key_column": "company",
    }
    rows = extract_rows(read_workbook(path), layout)
    assert [r["company"] for r in rows] == ["Acme"]
    assert rows[0]["__group__"] == "fund i"


def test_generated_workbook_bytes_are_deterministic(tmp_path):
    """The spec pins sha256(workbook); two generations must agree byte for byte
    (openpyxl stamps wall-clock dates into docProps and zip entries).

    NOTE what this does NOT establish: both generations run in ONE process, so
    they share a serialiser. That blind spot is the whole subject of
    `test_generated_workbook_bytes_are_writer_independent` below — this test was
    green for the entire period the pilot was machine-dependent.
    """
    import hashlib

    a, _, _ = _gen.build_one("L5_S1", tmp_path / "a")
    b, _, _ = _gen.build_one("L5_S1", tmp_path / "b")
    assert (
        hashlib.sha256(a.read_bytes()).hexdigest()
        == hashlib.sha256(b.read_bytes()).hexdigest()
    )


#: Generate one workbook and print its sha256. Run as a subprocess so the
#: `lxml`-blocking arm gets a clean interpreter -- openpyxl reads its `LXML`
#: flag at import time, so the choice cannot be changed once it is loaded.
_DIGEST_PROBE = """
import hashlib, importlib.util, sys
if {block_lxml!r}:
    class _Block:
        def find_spec(self, name, path=None, target=None):
            if name == "lxml" or name.startswith("lxml."):
                raise ImportError("lxml blocked by the writer-independence test")
            return None
    sys.meta_path.insert(0, _Block())
sys.path.insert(0, {root!r})
sys.path.insert(0, {pilot!r})
spec = importlib.util.spec_from_file_location("gs", {gen!r})
gs = importlib.util.module_from_spec(spec)
sys.modules["gs"] = gs
spec.loader.exec_module(gs)
from openpyxl import LXML
path, _, _ = gs.build_one("L2_S1", __import__("pathlib").Path({out!r}))
print(LXML, hashlib.sha256(path.read_bytes()).hexdigest())
"""


def test_generated_workbook_bytes_are_writer_independent(tmp_path):
    """The same seed must produce the same bytes with and without lxml.

    openpyxl selects its XML serialiser from `openpyxl.LXML` -- True when lxml
    is importable, False otherwise -- and the two writers disagree about
    attribute order, namespace-declaration placement and empty-element form
    (`<tag/>` vs `<tag />`). None of that carries meaning, but all of it changes
    the bytes, and the auditor's spec pins sha256(workbook). Measured 2026-09-03
    on openpyxl 3.1.5: seven of nine zip entries differed, so an honest bundle
    built without lxml was rejected PINNED_INPUT_MISMATCH. It passed in-house
    only because lxml happened to be installed as an unrelated transitive.

    `_generate_set._canonicalize_workbook` C14N-normalises each XML part, which
    makes the output a function of the document rather than of the writer.
    """
    import hashlib
    import subprocess

    lxml_spec = importlib.util.find_spec("lxml")
    if lxml_spec is None:
        # Both arms would use the pure-python writer and the comparison would be
        # a tautology. Say so rather than passing on nothing.
        pytest.skip(
            "lxml is not installed, so both arms would use the same serialiser "
            "and this test would be vacuous — install veriker[dev]"
        )

    digests = {}
    for label, block in (("with_lxml", False), ("without_lxml", True)):
        out = tmp_path / label
        code = _DIGEST_PROBE.format(
            block_lxml=block,
            root=str(_PKG_ROOT),
            pilot=str(_PILOT),
            gen=str(_PILOT / "_generate_set.py"),
            out=str(out),
        )
        proc = subprocess.run(
            [sys.executable, "-c", code],
            capture_output=True,
            text=True,
            cwd=str(_PKG_ROOT),
            env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        )
        assert proc.returncode == 0, f"{label} arm failed:\n{proc.stderr[-2000:]}"
        flag, digest = proc.stdout.strip().split()
        digests[label] = (flag, digest)

    # The arms must actually have run different serialisers, or the assertion
    # below proves nothing.
    assert digests["with_lxml"][0] == "True", (
        "the unblocked arm did not load lxml, so both arms used the same writer "
        f"and this test is vacuous: {digests}"
    )
    assert digests["without_lxml"][0] == "False", (
        f"blocking lxml did not take effect: {digests}"
    )

    assert digests["with_lxml"][1] == digests["without_lxml"][1], (
        "the workbook bytes depend on which XML serialiser openpyxl chose, so "
        f"the pinned sha256 is machine-dependent: {digests}"
    )


def test_a_content_edit_still_moves_the_workbook_digest(tmp_path):
    """Canonicalisation must not have cost the pin its teeth.

    The anti-tautology control for `_canonicalize_workbook`: C14N is only safe
    because it removes degrees of freedom that carry no content, so a change
    that DOES carry content has to survive it. Edits one numeric cell in the
    worksheet part and requires the digest to move.

    Without this, "canonicalisation cannot weaken the pin" would be an assertion
    in a docstring rather than a property under test — which is exactly what it
    was until the 2026-09-03 audit pointed out that the claim named a control
    that did not exist.
    """
    import hashlib
    import zipfile

    path, _, _ = _gen.build_one("L2_S1", tmp_path / "src")
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    zf = zipfile.ZipFile(path)
    entries = [(i, zf.read(i.filename)) for i in zf.infolist()]
    zf.close()

    edited_path = tmp_path / "edited.xlsx"
    edited = False
    with zipfile.ZipFile(edited_path, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet1.xml" and b"<v>" in data:
                i = data.index(b"<v>")
                j = data.index(b"</v>", i)
                data = data[: i + 3] + b"999999.0" + data[j:]
                edited = True
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 2, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, data)

    assert edited, "no numeric cell was found to edit — this control is vacuous"
    after = hashlib.sha256(edited_path.read_bytes()).hexdigest()
    assert before != after, (
        "a changed cell value left the workbook digest unchanged — the pinned "
        "input no longer binds the workbook's content"
    )


def test_generated_workbook_carries_no_toolchain_version(tmp_path):
    """The pinned digest must not be a digest of the build environment.

    openpyxl stamps its own version into `docProps/app.xml` ("Microsoft Excel
    Compatible / Openpyxl 3.1.5"). C14N normalises how a part is serialised,
    never what it contains, so it cannot remove that — and the extra declares
    `openpyxl>=3.1` unpinned. Measured 2026-09-03 before the fix: 3.1.5, 3.1.2
    and 3.0.10 produced three different sha256s and two PINNED_INPUT_MISMATCHes,
    with app.xml the only differing entry between 3.1.5 and 3.1.2.

    Checks the property rather than the symptom, so it also catches a future
    part that starts carrying a version string.
    """
    import re
    import zipfile

    path, _, _ = _gen.build_one("L2_S1", tmp_path)
    zf = zipfile.ZipFile(path)
    version_like = re.compile(rb"\d+\.\d+\.\d+")
    offenders = []
    for info in zf.infolist():
        raw = zf.read(info.filename)
        if b"Openpyxl" in raw or b"openpyxl" in raw:
            offenders.append((info.filename, "names openpyxl"))
        elif info.filename.startswith("docProps/") and version_like.search(raw):
            offenders.append((info.filename, "carries a dotted version"))
    assert not offenders, (
        f"these parts bake the build toolchain into the pinned digest: {offenders}"
    )


def test_canonicalisation_refuses_a_dtd_instead_of_inlining_it(tmp_path):
    """C14N must not launder a DTD past the verifier's byte-level guard.

    `sheet_query._guard_xml` refuses a workbook whose RAW bytes contain
    `<!DOCTYPE` or `<!ENTITY` — added 2026-09-02 after an adversarial pass got
    an entity expanded into a cell past a prefix-only scan. `ET.canonicalize`
    deletes the internal subset and inlines the expansion, so canonicalising
    such a part would leave the verifier nothing to find: the red-team lens on
    2026-09-03 built two workbooks differing only in an entity-expanded cell,
    canonicalised them to the SAME digest, and had the verifier accept one and
    refuse the other.

    So `_c14n` refuses instead of normalising. The generator never emits a DTD;
    this test is the control proving the refusal is wired, not the expectation
    that it fires in normal use.
    """
    entity_part = (
        b'<?xml version="1.0"?>\n'
        b'<!DOCTYPE worksheet [<!ENTITY nav "26.71">]>\n'
        b"<worksheet><v>&nav;</v></worksheet>"
    )
    with pytest.raises(ValueError, match="DTD/entity"):
        _gen._c14n(entity_part, "xl/worksheets/sheet1.xml")

    # And the guard is on the CONTENT, not the filename.
    with pytest.raises(ValueError, match="DTD/entity"):
        _gen._c14n(b'<!DOCTYPE Types><Types xmlns="urn:x"/>', "[Content_Types].xml")

    # Control: an ordinary part is still canonicalised rather than refused.
    assert _gen._c14n(b'<a x="1" ><b /></a>', "x.xml") == b'<a x="1"><b></b></a>'


def test_workbook_xml_parts_are_already_canonical(tmp_path):
    """Every XML part is in C14N form, so re-canonicalising is a no-op.

    The property behind the writer-independence guarantee: canonicalising ANY
    serialiser's output lands on these bytes, so two serialisers that emit the
    same document converge here. (Not the converse — openpyxl's own writers do
    not emit canonical bytes; that divergence is the thing being normalised
    away.) Cheaper than the subprocess test and it localises a regression to the
    offending part.
    """
    import io
    import zipfile
    from xml.etree import ElementTree as ET

    path, _, _ = _gen.build_one("L2_S1", tmp_path)
    zf = zipfile.ZipFile(path)
    checked = []
    for info in zf.infolist():
        if not info.filename.endswith((".xml", ".rels")):
            continue
        raw = zf.read(info.filename)
        out = io.StringIO()
        ET.canonicalize(xml_data=raw.decode("utf-8"), out=out)
        assert out.getvalue().encode("utf-8") == raw, (
            f"{info.filename} is not in canonical form, so its bytes still "
            "depend on which serialiser wrote it"
        )
        checked.append(info.filename)
    assert len(checked) >= 7, (
        f"expected the workbook to carry several XML parts, saw {checked} — "
        "a shrinking denominator would make this check vacuous"
    )


def test_harness_rounds_half_up_and_records_abstention():
    """MODERATE findings: float round() is half-even (85.945 -> 85.94, nine set
    questions sit on that boundary); an explicit null is an abstention, not an
    unparseable answer."""
    spec = importlib.util.spec_from_file_location(
        "finsheet_style_minimal__harness_run", _PILOT / "harness" / "run.py"
    )
    harness = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = harness
    spec.loader.exec_module(harness)

    assert harness.normalise_answer("number", 85.945) == 85.95
    assert harness.normalise_answer("number", "85.945") == 85.95
    assert harness.normalise_answer("number", "$1,234.565m") == 1234.57
    assert harness.normalise_answer("number", "NaN") is None
    assert harness.parse_model_json('{"answer": null, "derivation": {}}') == {
        "answer": None,
        "derivation": {},
    }
